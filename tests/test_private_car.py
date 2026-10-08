"""
Prove the PRIVATE CAR side of the quote matrix reaches the right verdicts -
offline, in seconds.

    venv\\Scripts\\python -m pytest tests/test_private_car.py -q

Same idea as tests/test_quote_matrix.py: every rule is shown a case it must
catch and a case it must leave alone. The car rules come from Saarthi's
pc-dont-know-number component and were checked on the live form on
2026-10-05 (Comprehensive up to 14 years, OD Only offered up to 4 but stuck
at 4, Third Party only from 15).

Nothing here touches the portal, an insurer, or any real notebook. The
browser tests use a local page only.
"""
from __future__ import annotations

import json
import sys
from datetime import date, timedelta
from pathlib import Path
from types import SimpleNamespace

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from core import (matrixexcel, matrixreport, quotechecks, quotenotes,  # noqa: E402
                  vehiclecatalog, vehiclenotes)
from core.quotecapture import Decline, PlanAnswer, QuoteCapture  # noqa: E402
from data import matrix  # noqa: E402
from data.matrix import CP, OD, OVER_90, TP  # noqa: E402
from pages import routes  # noqa: E402
from pages.quote_list import Quote  # noqa: E402

YEAR = matrix.this_year()
SWIFT = "MARUTI NEW SWIFT 1.2 ZXI AMT (1197 CC) (PETROL)"
WAGONR = "MARUTI WAGON R GREEN LXI (998 CC) (CNG)"
SCORPIO = "MAHINDRA AND MAHINDRA SCORPIO 2WD AC (2179 CC) (DIESEL)"
NEXON = "TATA NEXON EV XM (ELECTRIC)"
BANDS = {SWIFT: "1000-1500cc", WAGONR: "up to 1000cc", SCORPIO: "over 1500cc",
         NEXON: "electric"}
FACTS = {SWIFT: {"product": "car", "cc": 1197.0, "fuel": "PETROL"},
         WAGONR: {"product": "car", "cc": 998.0, "fuel": "CNG"},
         SCORPIO: {"product": "car", "cc": 2179.0, "fuel": "DIESEL"},
         NEXON: {"product": "car", "cc": None, "fuel": "ELECTRIC"}}


def opts(cars=(SWIFT,), insurers=("BAJAJ", "ICICI")):
    return matrix.options(list(cars), ["GJ-01 Ahmedabad", "MH-01 Mumbai"],
                          list(insurers), "car")


def car(**change):
    base = {"vehicle": SWIFT, "rto": "GJ-01 Ahmedabad", "year": str(YEAR - 3),
            "policy": CP, "previous": matrix.NOT_EXPIRED, "prev_type": CP,
            "prev_insurer": "BAJAJ", "ncb": "default", "claim": "No"}
    base.update(change)
    return base


def answer(code, premium=5119.0, od=1418.26, tp=3416.0, idv=432000.0, ncb=35.0,
           ncb_rs=496.39, pa=0.0, **kw):
    net = round(premium / 1.18, 2)
    base = dict(insurer=code, plan_id="1", ok=True, premium=premium, net=net,
                gst=round(premium - net, 2), od=od, tp=tp, pa_cover=pa, idv=idv,
                idv_min=idv * 0.85, idv_max=idv * 1.2, ncb_percent=ncb,
                ncb_discount=ncb_rs, years="1")
    base.update(kw)
    return PlanAnswer(**base)


def baseline(**change):
    return matrix.baseline(opts(), "car").with_(**change)


def run_checks(s, offers, declines=(), sent=None, vehicle=SWIFT):
    cards = [Quote(insurer=c, premium=a.premium, idv=None) for c, a in offers.items()]
    return quotechecks.check_journey(1, s, offers, list(offers.values()),
                                     list(declines), sent or {}, cards,
                                     {"BAJAJ": "BAJAJ", "ICICI": "ICICI"},
                                     FACTS.get(vehicle))


# ============================================================ the car form

def test_car_form_rules_are_not_the_bike_rules():
    ok = lambda **kw: matrix.normalise(car(**kw), "car")
    assert ok(year=str(YEAR - 14)) is not None                 # CP: up to 14 years
    assert ok(year=str(YEAR - 15)) is None
    assert ok(year=str(YEAR - 16), policy=TP) is not None      # TP only after that
    assert ok(policy=OD, year=str(YEAR - 4)) is not None       # offered at 4 ...
    assert ok(policy=OD, year=str(YEAR - 5)) is None           # ... not at 5
    # The same 15-year-old Comprehensive is fine for a bike (up to 25).
    bike = dict(car(year=str(YEAR - 15)), vehicle="HONDA ACTIVA")
    assert matrix.normalise(bike) is not None


def test_every_planned_car_journey_is_one_the_car_form_allows():
    plan, _ = matrix.plan(opts((SWIFT, WAGONR, SCORPIO, NEXON)), budget=40,
                          product="car", bands=BANDS)
    assert plan[0].kind == "baseline"
    assert plan[0].get("year") == str(YEAR - 3)               # 3, not the bike's 4
    for s in plan:
        again = matrix.normalise(dict(zip(matrix.DIMENSIONS, s.values)), "car",
                                 bands=BANDS)
        assert again is not None, s.key
        assert tuple(again[d] for d in matrix.DIMENSIONS) == s.values, s.key


def test_car_years_reach_past_the_comprehensive_limit():
    years = {matrix.age_of(y) for y in opts()["year"]}
    assert years == set(matrix.CAR_AGES)
    assert max(years) > matrix.RULES["car"].cp_max_age      # a TP-only car too


def test_a_learned_form_block_is_left_out_of_the_plan():
    blocked = frozenset({matrix.block_key(OD, 4)})
    plan, _ = matrix.plan(opts((SWIFT, WAGONR)), budget=60, product="car",
                          blocked=blocked)
    assert not any(s.get("policy") == OD and matrix.age_of(s.get("year")) == 4
                   for s in plan)
    without, _ = matrix.plan(opts((SWIFT, WAGONR)), budget=60, product="car")
    assert any(s.get("policy") == OD and matrix.age_of(s.get("year")) == 4
               for s in without)                             # tried when not learned


def test_electric_cars_are_never_planned_older_than_they_can_be():
    plan, _ = matrix.plan(opts((SWIFT, NEXON)), budget=60, product="car",
                          bands=BANDS)
    evs = [s for s in plan if s.get("vehicle") == NEXON]
    assert evs and all(matrix.age_of(s.get("year")) <= matrix.EV_MAX_AGE for s in evs)


def test_car_twins_use_car_ages_and_car_engine_bands():
    base = matrix.baseline(opts((SWIFT, WAGONR, SCORPIO)), "car")
    twins = {t.relation: t for t in matrix.twins(
        base, opts((SWIFT, WAGONR, SCORPIO)), BANDS, "car")}
    assert twins["age-idv"].get("year") == str(YEAR - 7)
    assert twins["cc-tp"].get("vehicle") == SCORPIO          # bigger than the Swift
    assert twins["cp-vs-od"].get("policy") == OD             # a 3-year-old car can


# ============================================================ the wire

class FakeRequest:
    def __init__(self, url, body, method="POST"):
        self.url, self.post_data, self.method = url, json.dumps(body), method
        self.failure = "net::ERR_FAILED"


class FakeResponse:
    def __init__(self, request, body, status=200):
        self.request, self.url, self.status = request, request.url, status
        self._body = json.dumps(body)

    def text(self):
        return self._body


CAR_API = "http://localhost:53339/api/Motor/privatecar/"
BIKE_API = "http://localhost:53339/api/Motor/TwoWheeler/"


def test_a_car_capture_reads_the_car_fan_out_and_nothing_else():
    c = QuoteCapture(segment="privatecar")
    c._on_response(FakeResponse(FakeRequest(CAR_API + "QualifiedCompany", {}), {
        "Response": {"QuotationNumber": "PIBLMTRPC1",
                     "QualifiedPlanList": [{"CompanyCode": "SHRIRAM", "PlanId": "497"},
                                           {"CompanyCode": "CHOLAMANDLAM",
                                            "PlanId": "271"}]}}))
    assert c.quotation_no == "PIBLMTRPC1" and not c.complete
    # A bike call in the same browser is not one of ours.
    bike = FakeRequest(BIKE_API + "SHRIRAM", {"CompanyCode": "SHRIRAM", "PlanId": 1})
    assert not c._is_plan_call(bike)
    # Car-only endpoints under the same prefix are not answers either.
    assert not c._is_plan_call(FakeRequest(CAR_API + "CompanySpecificQuotation",
                                           {"CompanyCode": "SHRIRAM", "PlanId": 1}))
    for code, item in (("SHRIRAM", {"CompanyCode": "SHRIRAM", "Status": "Success",
                                    "FinalPremium": 5119, "PremiumYear": 1}),
                       ("CHOLAMANDLAM", {"CompanyCode": "CHOLAMANDLAM", "Status": 0})):
        req = FakeRequest(CAR_API + code, {"CompanyCode": code, "PlanId": "1"})
        c._on_request(req)
        c._on_response(FakeResponse(req, {"Response": [item]}))
    assert c.complete
    assert list(c.best_offers()) == ["SHRIRAM"]
    assert [d.reason for d in c.declines()] == ["Status=0"]


def test_a_bike_capture_ignores_car_calls():
    c = QuoteCapture()
    assert not c._is_plan_call(FakeRequest(CAR_API + "SHRIRAM",
                                           {"CompanyCode": "SHRIRAM", "PlanId": 1}))
    with pytest.raises(ValueError):
        QuoteCapture.attach(SimpleNamespace(on=lambda *a: None), "car")  # not a segment


# ============================================================ the checks

def test_live_car_baseline_replays_to_the_same_verdicts():
    """Journey #1 of the first live car run, 2026-10-05: SHRIRAM priced the
    Swift (TP exactly the regulated Rs 3,416, NCB 25% -> 35%), CHOLAMANDLAM
    answered Status 0 with all zeros, ROYALSUNDRAM an empty list."""
    a = answer("SHRIRAM")
    a.net, a.gst, a.idv_min, a.idv_max = 4338.0, 781.0, 367200.0, 518400.0
    found = run_checks(baseline(), {"SHRIRAM": a}, declines=[
        Decline("CHOLAMANDLAM", "Status=0", "insurer"),
        Decline("ROYALSUNDRAM", "empty answer - the page shows nothing for this "
                                "insurer", "insurer")])
    got = {(f.severity, f.check, f.insurer) for f in found}
    assert got == {("LOOK", "no-reason", "CHOLAMANDLAM"),
                   ("LOOK", "vanished", "ROYALSUNDRAM")}, got


def test_third_party_must_be_the_regulated_price():
    # Swift, 1197cc: Rs 3,416. Planted: Rs 3,000.
    clean = run_checks(baseline(), {"SHRIRAM": answer("SHRIRAM")})
    assert "tp-rate" not in {f.check for f in clean}
    wrong = run_checks(baseline(), {"SHRIRAM": answer("SHRIRAM", tp=3000.0)})
    assert [(f.severity, f.insurer) for f in wrong if f.check == "tp-rate"] == \
        [("LOOK", "SHRIRAM")]
    # The PA cover inside the TP part is not counted against it.
    with_pa = run_checks(baseline(), {"SHRIRAM": answer("SHRIRAM", tp=3741.0, pa=325.0)})
    assert "tp-rate" not in {f.check for f in with_pa}


def test_regulated_price_knows_engine_bands_gas_kits_and_electric():
    assert quotechecks.regulated_tp("car", 998, "CNG") == {2094, 2154}
    assert quotechecks.regulated_tp("car", 1500, "PETROL") == {3416}
    assert quotechecks.regulated_tp("car", 2179, "DIESEL") == {7897}
    assert quotechecks.regulated_tp("car", None, "ELECTRIC") == set()
    assert quotechecks.regulated_tp("car", 1498, "PETROL HYBRID") == set()
    assert quotechecks.regulated_tp("bike", 110, "PETROL") == {714}
    assert quotechecks.regulated_tp("bike", 411, "PETROL") == {2804}
    # Electric: no check at all, rather than a wrong one.
    s = baseline(vehicle=NEXON)
    found = run_checks(s, {"SHRIRAM": answer("SHRIRAM", tp=1780.0)}, vehicle=NEXON)
    assert "tp-rate" not in {f.check for f in found}


def test_a_car_refused_as_a_two_wheeler_plate_is_named():
    tp = baseline(policy=TP, ncb=None, claim=None)
    refusal = Decline("UNITED", "User is not authorized to underwrite Private Car "
                      "as opted Registration number belongs to TwoWheeler", "insurer")
    found = run_checks(tp, {"SHRIRAM": answer("SHRIRAM", od=0.0, idv=0.0, ncb=0.0,
                                              ncb_rs=0.0)},
                       declines=[refusal],
                       sent={"IsThirdPartyOnly": True,
                             "VehicleDetails": {"RegistrationNumber": "GJ-01-AB-1111"}})
    hit = [f for f in found if f.check == "placeholder-number"]
    assert hit and "two-wheeler" in hit[0].title and "GJ-01-AB-1111" in hit[0].title


def test_bigger_car_engine_must_not_get_cheaper_third_party():
    tp = baseline(policy=TP, ncb=None, claim=None)
    found, _ = quotechecks.compare(
        [(1, tp, {"SHRIRAM": answer("SHRIRAM", od=0.0, tp=3416.0)}),
         (2, tp.with_(vehicle=SCORPIO), {"SHRIRAM": answer("SHRIRAM", od=0.0,
                                                          tp=3000.0)})], [], BANDS)
    assert [f.check for f in found] == ["cc-tp"]
    fine, state = quotechecks.compare(
        [(1, tp, {"SHRIRAM": answer("SHRIRAM", od=0.0, tp=3416.0)}),
         (2, tp.with_(vehicle=SCORPIO), {"SHRIRAM": answer("SHRIRAM", od=0.0,
                                                          tp=7897.0)})], [], BANDS)
    assert fine == [] and state["cc-tp"] == "held"


# ============================================================ the notebooks

@pytest.fixture
def notebooks(tmp_path, monkeypatch):
    monkeypatch.setattr(quotenotes, "FILES", {"bike": tmp_path / "bike.json",
                                              "car": tmp_path / "car.json"})
    monkeypatch.setattr(quotenotes, "NOTES_FILE", tmp_path / "bike.json")
    return tmp_path


def test_cars_learn_into_their_own_notebook(notebooks):
    quotenotes.use("car")
    assert quotenotes.product() == "car"
    quotenotes.record_journey(baseline(), {"SHRIRAM": answer("SHRIRAM")}, [], {}, "Q")
    assert (notebooks / "car.json").exists() and not (notebooks / "bike.json").exists()
    quotenotes.use("bike")
    assert quotenotes.history_pairs() == {}                  # bikes see none of it


def test_a_form_block_is_remembered_for_a_fortnight_then_rechecked(notebooks):
    quotenotes.use("car")
    key = matrix.block_key(OD, 4)
    quotenotes.record_block(key, "TP Policy Expiry Date cannot be valid")
    assert quotenotes.blocked() == {key}
    later = date.today() + timedelta(days=quotenotes.BLOCK_RECHECK_DAYS)
    assert quotenotes.blocked(later) == frozenset()          # time to look again
    quotenotes.record_run(1, [], {}, [], {})
    assert "OD Only, 4 years old" in quotenotes.report()     # and it is reported
    quotenotes.record_unblock(key)
    assert quotenotes.blocked() == frozenset()


def test_suspects_that_always_came_together_are_one_rule(notebooks):
    """Live, 2026-10-05: CHOLAMANDLAM refused every Swift journey - all in
    GJ-01, all 2023, all with BAJAJ as previous insurer - and quoted the
    others. That is one rule with four suspects, not four rules."""
    quotenotes.use("car")
    refusal = Decline("CHOLAMANDLAM", "Status=0", "insurer")
    swift = baseline()
    for change in ({}, {"ncb": "0%"}, {"claim": "Yes", "ncb": None}):
        quotenotes.record_journey(swift.with_(**change), {}, [refusal],
                                  {"CHOLAMANDLAM": "unknown"}, "Q")
    other = swift.with_(vehicle=WAGONR, rto="MH-01 Mumbai", year=str(YEAR - 1),
                        prev_insurer="ICICI", policy=OD)
    quotenotes.record_journey(other, {"CHOLAMANDLAM": answer("CHOLAMANDLAM")}, [],
                              {}, "Q")
    learned = [t for c, t in quotenotes.rules("CHOLAMANDLAM")]
    assert len(learned) == 1, learned
    assert learned[0].startswith("refused all 3 journeys that had")
    for suspect in ("RTO = GJ-01 Ahmedabad", "previous insurer = BAJAJ",
                    f"vehicle = {SWIFT}"):
        assert suspect in learned[0]
    assert "--product car" in learned[0]


def test_an_insurer_outage_is_never_called_our_bug():
    bajaj_down = Decline("BAJAJ", 'Error in Service ~~ I/O error on POST request for '
                         '"https://htauth.preprod.bajajgeneral.com/...": '
                         'javax.net.ssl.SSLHandshakeException', "insurer")
    found = run_checks(baseline(policy=TP, ncb=None, claim=None),
                       {"SHRIRAM": answer("SHRIRAM", od=0.0, idv=0.0, ncb=0.0,
                                          ncb_rs=0.0)}, declines=[bajaj_down])
    assert not [f for f in found if f.insurer == "BAJAJ" and f.check == "our-code"]


def test_car_plates_are_kept_apart_from_bike_plates(tmp_path, monkeypatch):
    monkeypatch.setattr(vehiclenotes, "NOTES_FILE", tmp_path / "plates.json")
    vehiclenotes._save({"numbers": {}})
    vehiclenotes.record("ZUNO", "GJ01YX0850", True, "")            # a bike plate
    vehiclenotes.record("NATIONAL", "GJ01KC7427", False,
                        "Vehicle sub class is not matching with vehicle "
                        "registration number data.")                # not a bike
    plate, why = vehiclenotes.choose("GJ-01 Ahmedabad", "SHRIRAM", product="car")
    assert plate == "GJ01KC7427" and "sub class" in why             # the hint
    vehiclenotes.record("SHRIRAM", plate, True, "", product="car")
    plate, why = vehiclenotes.choose("GJ-01 Ahmedabad", "SHRIRAM", product="car")
    assert plate == "GJ01KC7427" and why.startswith("accepted before")
    # The bike pool is untouched and never offers the car's plate.
    assert vehiclenotes.choose("GJ-01 Ahmedabad", "ZUNO")[0] == "GJ01YX0850"
    assert vehiclenotes.known("car").startswith("1 proven")
    assert vehiclenotes.product_of("MTRPC") == "car"
    assert vehiclenotes.product_of("MTRTW") == "bike"
    assert routes.product_of("http://localhost:4200/private-car/proposal") == "car"
    assert routes.product_of("http://localhost:4200/two-wheeler/proposal") == "bike"


# ============================================================ picking cars

ROWS = [
    {"make": "MARUTI", "model": "NEW SWIFT", "variant": "1.2 ZXI AMT (1197 CC) (PETROL)",
     "fuel": "PETROL", "cc": 1197.0, "band": "1000-1500cc", "top": True},
    {"make": "MARUTI", "model": "WAGON R", "variant": "LXI (998 CC) (PETROL)",
     "fuel": "PETROL", "cc": 998.0, "band": "up to 1000cc", "top": True},
    {"make": "MARUTI", "model": "WAGON R", "variant": "GREEN LXI (998 CC) (CNG)",
     "fuel": "CNG", "cc": 998.0, "band": "up to 1000cc", "top": False},
    {"make": "MAHINDRA AND MAHINDRA", "model": "SCORPIO", "variant": "S4 (2179 CC) (PETROL)",
     "fuel": "PETROL", "cc": 2179.0, "band": "over 1500cc", "top": True},
    {"make": "MAHINDRA AND MAHINDRA", "model": "SCORPIO", "variant": "S4 (2179 CC) (DIESEL)",
     "fuel": "DIESEL", "cc": 2179.0, "band": "over 1500cc", "top": False},
    {"make": "TATA", "model": "NEXON EV", "variant": "XM (ELECTRIC)",
     "fuel": "ELECTRIC", "cc": None, "band": "electric", "top": False},
    {"make": "TATA", "model": "NEXON EV", "variant": "XZ (ELECTRIC)",
     "fuel": "ELECTRIC", "cc": None, "band": "electric", "top": False},
]


def test_cars_are_picked_one_per_band_with_a_spread_of_fuels():
    picked = vehiclecatalog.pick({"vehicles": ROWS}, 4, product="car")
    assert picked[0][0] == vehiclecatalog.PROVEN_CAR
    fuels = [vehiclecatalog.row_for(v, {"vehicles": ROWS}).get("fuel") for v, _ in picked]
    assert fuels == ["PETROL", "CNG", "DIESEL", "ELECTRIC"]
    assert [band for _, band in picked] == ["1000-1500cc", "up to 1000cc",
                                            "over 1500cc", "electric"]


def test_a_vehicle_named_the_notebooks_way_is_really_avoided():
    """The notebook writes 'TATA NEXON EV XM (ELECTRIC)'; the old picker only
    compared 'TATA|NEXON EV|XM (ELECTRIC)', so a refused vehicle came back."""
    picked = vehiclecatalog.pick({"vehicles": ROWS}, 4, product="car",
                                 avoid={"TATA NEXON EV XM (ELECTRIC)"},
                                 tried={"MAHINDRA AND MAHINDRA SCORPIO S4 (2179 CC) (DIESEL)"})
    labels = [vehiclecatalog.label(v) for v, _ in picked]
    assert "TATA NEXON EV XZ (ELECTRIC)" in labels           # the other EV instead
    assert "MAHINDRA AND MAHINDRA SCORPIO S4 (2179 CC) (PETROL)" in labels  # untried first


def test_cars_hidden_behind_a_same_named_make_are_found_and_never_picked():
    """The live car list, 2026-10-05: TATA twice (110, 391), the Nexon EV
    under the second - which the form, keeping the first 'TATA', never shows."""
    raw = {
        "makes": [{"Id": "110", "Name": "TATA"}, {"Id": "92", "Name": "MARUTI"},
                  {"Id": "391", "Name": "TATA "}],
        "models": [{"Id": "1024", "Name": "NEXON", "MakeId": "110"},
                   {"Id": "5313", "Name": "NEW SWIFT", "MakeId": "92"},
                   {"Id": "5372", "Name": "NEXON", "MakeId": "391"}],
        "variants": [{"VariantName": "1.2 XM (1198 CC)", "FuelType": "PETROL",
                      "ModelId": "1024"},
                     {"VariantName": "1.2 ZXI AMT (1197 CC)", "FuelType": "PETROL",
                      "ModelId": "5313"},
                     {"VariantName": "EV XM", "FuelType": "ELECTRIC", "ModelId": "5372"}],
    }
    cat = vehiclecatalog.build(raw, "test", "car")
    hidden = [r for r in cat["vehicles"] if r["hidden_by"]]
    assert [r["variant"] for r in hidden] == ["EV XM (ELECTRIC)"]
    [line] = vehiclecatalog.hidden(cat)
    assert "TATA is listed 2 times (ids 110, 391)" in line and "NEXON" in line
    picked = vehiclecatalog.pick(cat, 4, product="car")
    assert all(v.variant != "EV XM (ELECTRIC)" for v, _ in picked)


# ============================================================ the spreadsheet

def test_the_car_sheet_says_success_or_failure_with_company_and_segment(tmp_path):
    s = baseline()
    ok = SimpleNamespace(n=1, scenario=s, status="ok", note="", quotation="PIBLMTRPC1",
                         offers={"SHRIRAM": answer("SHRIRAM")}, seconds=44,
                         declines=[Decline("CHOLAMANDLAM", "Status=0", "insurer")],
                         kinds={"CHOLAMANDLAM": "unknown"}, silent=[])
    stuck = SimpleNamespace(n=2, scenario=s.with_(policy=OD, year=str(YEAR - 4)),
                            status="blocked", quotation="", offers={}, declines=[],
                            kinds={}, silent=[], seconds=30,
                            note="screen 3: Proceed stays disabled")
    skipped = SimpleNamespace(n=3, scenario=stuck.scenario, status="skipped",
                              quotation="", offers={}, declines=[], kinds={},
                              silent=[], seconds=0, note="not run: #2 showed it")
    bug = quotechecks.Finding("DEFECT", "form-block", "FORM", "the form blocks it",
                              journeys=(2,))
    rows = matrixexcel.rows_for([ok, stuck, skipped], [bug], {}, "car",
                                {"SHRIRAM": "26", "CHOLAMANDLAM": "14"})
    by = {(r["n"], r["company"]): r for r in rows}
    assert by[(1, "SHRIRAM")]["status"] == "Success"
    assert by[(1, "SHRIRAM")]["company_id"] == "26"
    assert by[(1, "SHRIRAM")]["sub_product"] == "MTRPC - Private Car"
    assert by[(1, "SHRIRAM")]["segment"] == "1 - Comprehensive"
    assert by[(1, "CHOLAMANDLAM")]["status"] == "Failure"
    assert by[(1, "CHOLAMANDLAM")]["reason"] == "Status=0"
    assert by[(2, "-")]["status"] == "Failure" and "Bug:" in by[(2, "-")]["notes"]
    assert by[(3, "-")]["status"] == "Not run"
    words = {r["status"] for r in rows}
    assert words <= {"Success", "Failure", "Not run"}           # no jargon
    paths = matrixreport.write(tmp_path, [ok, stuck, skipped], ["SHRIRAM",
                               "CHOLAMANDLAM"], [bug], {}, [], [],
                               {"this_run": 1, "wanted": 2, "with_this_run": 1}, {},
                               "car", {"SHRIRAM": "26"})
    assert "Private Car" in paths["html"].read_text(encoding="utf-8")
    if "xlsx" in paths:
        from openpyxl import load_workbook
        sheet = load_workbook(paths["xlsx"])["Results"]
        cells = {c.value for row in sheet.iter_rows() for c in row}
        assert {"Success", "Failure", "Company ID", "Sub Product", "Segment"} <= cells


# ======================================================= in a real browser

def _chromium():
    try:
        from playwright.sync_api import sync_playwright
        pw = sync_playwright().start()
        return pw, pw.chromium.launch()
    except Exception as exc:                     # noqa: BLE001
        pytest.skip(f"no browser available: {exc}")


# What the live car form looked like on 2026-10-05 for a 4-year-old car on
# OD Only: the TP expiry date pre-filled in the past, marked invalid, and a
# grey Proceed - with no message anywhere.
STUCK = """
<form>
  <input formcontrolname="tpPolicyInsurer" class="ng-valid" value="BAJAJ ALLIANZ">
  <input formcontrolname="tpPolicyExpDate" class="ng-invalid" value="05 Oct 2025">
  <input formcontrolname="hiddenTab" class="ng-invalid" value="x" style="display:none">
  <button type="button" disabled>Proceed</button>
</form>
"""


def test_a_grey_proceed_is_explained_by_the_field_that_holds_it(tmp_path):
    from core import ui
    import run_quote_matrix as runner
    pw, chromium = _chromium()
    try:
        page = chromium.new_page()
        page.set_content(STUCK)
        invalid = ui.proceed_blocked(page, settle_ms=100)
        assert invalid == [("tpPolicyExpDate", "05 Oct 2025")]   # hidden tab ignored
        j = runner.Journey(6, baseline(policy=OD, year=str(YEAR - 4)))
        runner._blocked(j, page, tmp_path, "screen 3", invalid, "car")
        assert (tmp_path / "journey-06-blocked.png").exists()    # the evidence
        assert j.status == "blocked"
        assert "TP Policy Expiry Date" in j.note and "4-year-old car" in j.note
        assert [(f.severity, f.check) for f in j.findings] == [("DEFECT", "form-block")]
        page.set_content(STUCK.replace(" disabled", ""))
        assert ui.proceed_blocked(page, settle_ms=100) is None
    finally:
        chromium.close()
        pw.stop()
