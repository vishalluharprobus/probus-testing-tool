"""
The insurer lab's spreadsheet: one row per scenario, coloured so it reads at a glance.

    Results   #, company id, sub product, segment, scenario, then one column
              per stage - Quote, KYC, Company Specific, Proposal, Payment -
              each Success (green), Error (red), Not reached / Not run (grey);
              then Status (Success or Failure), the reason, notes and figures
    Summary   what was run, how each stage went, and the failure reasons
              grouped, most common first

Needs openpyxl. Without it write() returns False and the caller keeps the
plain CSV.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:                 # pragma: no cover - reported by write()
    Workbook = None

from core.labjourney import ERROR, NOT_REACHED, NOT_RUN, STAGES, SUCCESS


@dataclass(frozen=True)
class Column:
    heading: str
    key: str
    width: int
    kind: str = "text"     # text | wrap | int | money | pct | secs | stage | status


COLUMNS = (
    Column("#", "n", 5, "int"),
    Column("Company ID", "company_id", 11),
    Column("Company", "company", 13),
    Column("Sub Product", "sub_product", 19),
    Column("Segment", "segment", 18),
    Column("Scenario", "scenario", 50, "wrap"),
    *(Column(name, key, 17 if key == "company_specific" else 12, "stage")
      for key, name in STAGES),
    Column("Status", "status", 11, "status"),
    Column("Reason", "reason", 60, "wrap"),
    Column("Notes", "notes", 42, "wrap"),
    Column("Premium (Rs)", "premium", 13, "money"),
    Column("OD (Rs)", "od", 11, "money"),
    Column("TP (Rs)", "tp", 11, "money"),
    Column("IDV (Rs)", "idv", 12, "money"),
    Column("NCB %", "ncb", 8, "pct"),
    Column("Quotation No.", "quotation", 30),
    Column("Proposal No.", "proposal_no", 22),
    Column("Time (s)", "seconds", 9, "secs"),
)
FROZEN = "G6"            # header rows and #..Scenario stay put while scrolling
HEADER_ROW = 5

# Excel's own "Good / Bad / Neutral" colours, so they look familiar.
GREEN, GREEN_TEXT = "C6EFCE", "006100"
RED, RED_TEXT = "FFC7CE", "9C0006"
GREY, GREY_TEXT = "EDEDED", "7F7F7F"
AMBER, AMBER_TEXT = "FFEB9C", "7F6000"
NAVY, BAND, LINE = "1F3864", "F5F8FC", "D9DEE7"

STATUS_LOOK = {SUCCESS: (GREEN, GREEN_TEXT), "Failure": (RED, RED_TEXT),
               ERROR: (RED, RED_TEXT), NOT_REACHED: (GREY, GREY_TEXT),
               NOT_RUN: (GREY, GREY_TEXT)}


def write(path: Path, title: str, subtitle: str, rows: list[dict],
          facts: list[tuple[str, str]]) -> bool:
    """Write the workbook. False when openpyxl is missing."""
    if Workbook is None:
        return False
    book = Workbook()
    _results(book.active, title, subtitle, rows)
    _summary(book.create_sheet("Summary"), title, rows, facts)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return True


# ================================================================ Results

def _results(ws, title: str, subtitle: str, rows: list[dict]) -> None:
    ws.title = "Results"
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="595959")

    # A legend, in the colours it explains.
    ws["A3"] = "Key:"
    ws["A3"].font = Font(bold=True, color="595959")
    for col, word in zip("BCDE", (SUCCESS, ERROR, NOT_REACHED, NOT_RUN)):
        cell = ws[f"{col}3"]
        cell.value = word
        _paint(cell, *STATUS_LOOK[word], bold=True)
        cell.alignment = Alignment(horizontal="center")
    ws["F3"] = ("Every row starts from the standard quote and changes only what "
                "its Scenario says. A stage after an Error is Not reached.")
    ws["F3"].font = Font(italic=True, color="595959")

    thin = Side(style="thin", color=LINE)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=i, value=col.heading)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(i)].width = col.width
    ws.row_dimensions[HEADER_ROW].height = 30

    for r, row in enumerate(rows, start=HEADER_ROW + 1):
        band = PatternFill("solid", fgColor=BAND) if r % 2 == 0 else None
        for i, col in enumerate(COLUMNS, start=1):
            value = row.get(col.key)
            cell = ws.cell(row=r, column=i, value=None if value in ("", None) else value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=col.kind == "wrap",
                                       horizontal="center" if col.kind in
                                       ("stage", "status", "int", "pct") else None)
            if band is not None:
                cell.fill = band
            if col.kind == "money":
                cell.number_format = "#,##0"
            elif col.kind == "secs":
                cell.number_format = "0.0"
            elif col.kind in ("stage", "status") and value in STATUS_LOOK:
                _paint(cell, *STATUS_LOOK[value], bold=col.kind == "status")
            elif col.key == "reason" and value:
                cell.font = Font(color=RED_TEXT)
            elif col.key == "notes" and value:
                bug = "Bug:" in str(value)
                _paint(cell, AMBER, RED_TEXT if bug else AMBER_TEXT, bold=bug)

    last = HEADER_ROW + max(1, len(rows))
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(COLUMNS))}{last}"
    ws.freeze_panes = FROZEN


# ================================================================ Summary

def _summary(ws, title: str, rows: list[dict], facts) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 70
    for col in "CEF":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["D"].width = 22
    ws["A1"] = f"{title} - summary"
    ws["A1"].font = Font(size=16, bold=True, color=NAVY)

    r = 3
    for label, value in facts:
        bug = label.lower().startswith("portal bug")
        ws.cell(row=r, column=1, value=label).font = Font(bold=True,
                                                          color=RED_TEXT if bug else None)
        cell = ws.cell(row=r, column=2, value=value)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if bug:
            _paint(cell, RED, RED_TEXT)
        r += 1

    # How each stage went.
    r += 1
    ws.cell(row=r, column=1, value="How each stage went").font = Font(size=13, bold=True,
                                                                     color=NAVY)
    r += 1
    heads = ("Stage", "", SUCCESS, ERROR, NOT_REACHED, NOT_RUN)
    for i, head in enumerate(heads, start=1):
        cell = ws.cell(row=r, column=i, value=head or None)
        if head in STATUS_LOOK:
            _paint(cell, *STATUS_LOOK[head], bold=True)
            cell.alignment = Alignment(horizontal="center")
        else:
            cell.font = Font(bold=True)
    for key, name in STAGES:
        r += 1
        ws.cell(row=r, column=1, value=name).font = Font(bold=True)
        for i, word in enumerate(heads[2:], start=3):
            count = sum(1 for row in rows if row.get(key) == word)
            cell = ws.cell(row=r, column=i, value=count)
            cell.alignment = Alignment(horizontal="center")
            if count and word == ERROR:
                cell.font = Font(bold=True, color=RED_TEXT)

    # Why things failed, grouped: the same message with different numbers
    # in it is one reason.
    r += 2
    ws.cell(row=r, column=1, value="Why rows failed").font = Font(size=13, bold=True,
                                                                 color=NAVY)
    r += 1
    for i, head in enumerate(("Stage", "Reason", "Rows", "Row numbers"), start=1):
        ws.cell(row=r, column=i, value=head).font = Font(bold=True)
    groups: dict[tuple[str, str], list] = {}
    for row in rows:
        if row.get("status") != "Failure":
            continue
        stage, _, why = str(row.get("reason") or "").partition(": ")
        shape = re.sub(r"\d[\d,.]*", "#", why)
        groups.setdefault((stage, shape), []).append(row)
    if not groups:
        r += 1
        ws.cell(row=r, column=1, value="Nothing failed.").font = Font(color=GREEN_TEXT,
                                                                     bold=True)
    for (stage, _), members in sorted(groups.items(), key=lambda g: -len(g[1])):
        r += 1
        why = str(members[0].get("reason") or "").partition(": ")[2]
        ws.cell(row=r, column=1, value=stage).font = Font(bold=True, color=RED_TEXT)
        cell = ws.cell(row=r, column=2, value=why)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=3, value=len(members)).alignment = Alignment(
            horizontal="center", vertical="top")
        numbers = ", ".join(f"#{m.get('n')}" for m in members[:30])
        cell = ws.cell(row=r, column=4, value=numbers + (" ..." if len(members) > 30 else ""))
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _paint(cell, fill: str, text: str, bold: bool = False) -> None:
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(color=text, bold=bold)
